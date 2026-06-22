from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Plasmit-HMS-ICU-Command-Center-Executive-Use-Case-Overview.docx"

NAVY = "0F172A"
BLUE = "0284C7"
BLUE_DARK = "0369A1"
BLUE_LIGHT = "E0F2FE"
SLATE = "475569"
SLATE_LIGHT = "F1F5F9"
WHITE = "FFFFFF"
GREEN = "0F766E"
AMBER = "B45309"


MODULES = [
    (
        "Command",
        "Central visibility and prioritization for the entire ICU service.",
        [
            ("Command Center", "ICU Head, Duty Doctor, Hospital Admin", "Review patient risk, ventilation, fluid balance, medication alerts and open events in one command view."),
            ("Executive Dashboard", "Management, Hospital Admin, ICU Head", "Compare ICU units by occupancy, critical load, quality, devices, transfers and accountable leadership."),
            ("Notifications & Tasks", "ICU Head, Doctors, Nurses", "Prioritize clinical notifications and operational tasks by severity, SLA, status and assigned team."),
        ],
    ),
    (
        "Patients",
        "Patient discovery, bed visibility and admission-to-discharge flow.",
        [
            ("Patient Search", "Clinical and Administrative Teams", "Find ICU patients by identity, bed, unit, doctor, risk, ventilation and current status."),
            ("Smart Bed View", "ICU Head, Unit Nurse, Bed Manager", "View each bed with patient, care team, device readiness and patient-flow context."),
            ("Admissions", "Admission Desk, Unit Nurse, ICU Doctor", "Validate ICU need, select an appropriate bed and coordinate safe patient receipt."),
            ("Discharges", "ICU Doctor, Unit Nurse, Bed Manager", "Track readiness, clearances, destination, handover, transport and ICU bed release."),
        ],
    ),
    (
        "Critical Care",
        "Daily ICU operations, clinical risk management and coordinated action.",
        [
            ("ICU Operations", "ICU Head, Head Nurse, Bed Manager", "Monitor capacity, bed readiness, staffing, transfers and operational blockers."),
            ("Device Monitoring", "Biomedical Team, Head Nurse", "Monitor bedside devices, connectivity, last data, issues and corrective ownership."),
            ("Clinical Alerts", "Duty Doctor, Head Nurse, Ward Nurse", "Review clinically significant triggers with severity, source, SLA and required action."),
            ("ICU Rounds", "Intensivist, ICU Doctors, Multidisciplinary Team", "Conduct patient-wise rounds, review results and record the daily clinical plan."),
            ("Escalation Center", "ICU Head, Duty Doctor, Head Nurse", "Coordinate unresolved critical issues by priority, trigger, source, Assigned To, SLA and review status."),
        ],
    ),
    (
        "Clinical Workspace",
        "A unified patient-level workspace for review, documentation and care coordination.",
        [
            ("Patient Overview", "Doctors, Nurses, Allied Clinical Team", "Review the selected patient's current status, care team, observations, medication and active issues."),
            ("Progress Notes", "Doctors, Nurses, Allied Clinical Team", "Record structured clinical progress, procedures, events and follow-up notes."),
            ("Orders & Care Plans", "Doctors, Nurses", "Create and acknowledge medication, monitoring, procedure and nursing care instructions."),
            ("Family Communication", "Doctor, Unit Nurse, Family Coordinator", "Document family updates, consent discussions, questions and follow-up commitments."),
        ],
    ),
    (
        "Nursing",
        "Shift execution, bedside documentation, medication safety and handover continuity.",
        [
            ("Nursing Station", "Head Nurse, Unit Nurse, Ward Nurse", "Manage patient assignments, workload, due work and shift-level clinical priorities."),
            ("Nurse Entry", "Ward Nurse", "Capture bedside observations, vitals, oxygen support, neurological status, pain and urine output."),
            ("Medication Administration", "Ward Nurse, Pharmacy, Unit Nurse", "Execute eMAR workflows for due, administered, held, missed and high-alert medicines."),
            ("Patient Medication Chart", "Doctors, Nurses, Clinical Reviewers", "Review the patient medication history, timing, changes and administration status in read-only chart form."),
            ("Shift Handover", "Outgoing Nurse, Incoming Nurse, Unit Nurse", "Transfer pending work, clinical watch points and accountability between nursing shifts."),
            ("Tasks & Assessments", "Ward Nurse, Head Nurse", "Create, assign and complete nursing tasks while recording assessment and follow-up needs."),
        ],
    ),
    (
        "Diagnostics",
        "Consolidated diagnostics visibility and structured report intake.",
        [
            ("Diagnostics Hub", "Doctors, Lab, Radiology, Nurses", "Review laboratory, microbiology, imaging and specialty diagnostic results with critical-value status."),
            ("Report Upload & Extract", "Diagnostics Staff, Lab Technician", "Upload diagnostic reports, extract structured values and validate results before clinical use."),
        ],
    ),
    (
        "Tele ICU",
        "Remote specialist access for timely review of high-risk and resource-limited cases.",
        [
            ("Remote Command Center", "Remote Intensivist, ICU Head", "View remote-review readiness, patient risk, local team status and response priorities."),
            ("Remote Consultations", "Local Doctor, Remote Specialist", "Manage consultation requests, specialty routing, supporting documents and recommendations."),
            ("Escalated Cases", "Remote Intensivist, Local ICU Team", "Track high-priority cases from escalation trigger through action, SLA and outcome."),
        ],
    ),
    (
        "Device Operations",
        "Clinical engineering oversight for connected ICU equipment and data continuity.",
        [
            ("Edge Device Management", "Biomedical Team, Clinical Engineering", "Manage monitor, ventilator, pump and gateway inventory with operational status."),
            ("Device Mapping", "Biomedical Team, Unit Nurse", "Map devices to the correct ICU bed and patient context."),
            ("Connectivity Dashboard", "Biomedical Team, IT Operations", "Monitor connectivity, downtime and recovery actions across ICU devices and gateways."),
            ("Signal Health", "Biomedical Team, IT Operations", "Identify weak, delayed or missing signals before clinical data continuity is affected."),
        ],
    ),
    (
        "Clinical Intelligence",
        "Explainable risk identification to support earlier clinical intervention.",
        [
            ("Patient Risk Center", "ICU Doctor, Head Nurse, Quality Team", "Prioritize patients using combined signals from vitals, ventilation, infection, medication, devices and tasks."),
            ("Early Warning Scores", "Duty Doctor, Ward Nurse", "Track score trends, contributing observations and escalation thresholds."),
        ],
    ),
    (
        "Analytics",
        "Management insight for performance, quality, technology and adoption.",
        [
            ("Operational Analytics", "ICU Leadership, Hospital Admin", "Review occupancy, demand, length of stay, response times and workload."),
            ("Clinical Analytics", "Clinical Leadership, Quality Team", "Review safety, infection, medication, ventilation and outcome indicators."),
            ("Device Analytics", "Biomedical Leadership, IT", "Measure utilization, downtime, recurring issues and service performance."),
            ("Pilot Outcome Dashboard", "Management, Program Leadership", "Assess pilot KPIs, response improvement, workflow closure and outcome trends."),
            ("Adoption Analytics", "Hospital Admin, Training, Product Team", "Measure role-wise usage, workflow adherence and training needs."),
        ],
    ),
    (
        "Administration",
        "Controlled access, configurable rules and complete operational accountability.",
        [
            ("Users & Roles", "Hospital Admin, Security Admin", "Control role-based access for clinical, operational, remote and support users."),
            ("Configuration", "Hospital Admin, ICU Governance Team", "Configure units, beds, thresholds, escalation rules, medication timing and device settings."),
            ("Audit Logs", "Hospital Admin, Quality, Information Security", "Review who viewed or changed data, what changed and when the action occurred."),
        ],
    ),
]


PERSONAS = [
    ("Senior Management / Hospital Admin", "Unit performance, capacity, quality, adoption, governance and audit oversight."),
    ("ICU Head / Intensivist", "Clinical prioritization, rounds, escalations, care decisions and unit-level command."),
    ("Duty / Admitting / Consulting Doctor", "Patient review, orders, notes, results, urgent action and discharge decisions."),
    ("Head / Unit Nurse", "Staff allocation, workload supervision, escalation and handover assurance."),
    ("Ward Nurse", "Bedside observations, medication, tasks, assessments and shift handover."),
    ("Diagnostics Team", "Report intake, structured result review and critical-result communication."),
    ("Biomedical / IT Operations", "Device mapping, connectivity, signal health and service recovery."),
    ("Remote Intensivist", "Tele-ICU consultation, remote review and escalated-case support."),
]


FLOW = [
    ("1", "Admission & Bed Allocation", "Confirm ICU need, bed capability, patient identity and receiving team."),
    ("2", "Continuous Monitoring", "Capture bedside observations and receive connected-device data."),
    ("3", "Clinical Review & Orders", "Review trends, diagnostics and risk; document rounds, orders and care plans."),
    ("4", "Nursing Execution", "Administer medication, complete care tasks and document response."),
    ("5", "Diagnostics & Device Assurance", "Validate results, critical values, device mapping and data continuity."),
    ("6", "Alert, Escalation & Tele ICU", "Route unresolved risk to the right person/team within the defined SLA."),
    ("7", "Handover, Discharge & Learning", "Transfer accountability, release the bed and feed analytics/audit."),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def set_run(run, size=None, bold=None, color=None, font="Aptos"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_heading(doc, text, level=1, subtitle=None):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[f"Heading {level}"]
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run(run, size=18 if level == 1 else 13, bold=True, color=BLUE_DARK if level == 1 else NAVY)
    if subtitle:
        sub = doc.add_paragraph(subtitle)
        sub.paragraph_format.space_after = Pt(7)
        set_run(sub.add_run(), size=9, color=SLATE)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), size=9.5, bold=True, color=NAVY)
        set_run(p.add_run(text[len(bold_prefix):]), size=9.5, color=SLATE)
    else:
        set_run(p.add_run(text), size=9.5, color=SLATE)
    return p


def add_bullet(doc, text, color=SLATE):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    set_run(p.add_run(text), size=9.2, color=color)
    return p


def style_header_row(row, font_size=8.3):
    set_repeat_table_header(row)
    for cell in row.cells:
        set_cell_shading(cell, BLUE_DARK)
        set_cell_margins(cell, top=85, bottom=85)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_run(run, size=font_size, bold=True, color=WHITE)


def add_module_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(1.85)
    table.columns[2].width = Inches(4.15)
    table.rows[0].cells[0].text = "PAGE"
    table.rows[0].cells[1].text = "PRIMARY USERS"
    table.rows[0].cells[2].text = "HIGH-LEVEL USE CASE"
    style_header_row(table.rows[0])
    for index, (page, users, use_case) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = page
        cells[1].text = users
        cells[2].text = use_case
        fill = WHITE if index % 2 == 0 else "F8FAFC"
        for cell in cells:
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": "CBD5E1"})
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run(run, size=8.4, color=SLATE)
        for run in cells[0].paragraphs[0].runs:
            set_run(run, size=8.5, bold=True, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_two_column_cards(doc, cards):
    table = doc.add_table(rows=(len(cards) + 1) // 2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.7)
    table.columns[1].width = Inches(3.7)
    for idx, (title, detail) in enumerate(cards):
        cell = table.cell(idx // 2, idx % 2)
        set_cell_shading(cell, "F8FAFC")
        set_cell_border(cell, top={"val": "single", "sz": "5", "color": "CBD5E1"}, bottom={"val": "single", "sz": "5", "color": "CBD5E1"}, start={"val": "single", "sz": "14", "color": BLUE}, end={"val": "single", "sz": "5", "color": "CBD5E1"})
        set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(3)
        set_run(p1.add_run(title), size=9.2, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(detail), size=8.4, color=SLATE)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(SLATE)
    normal.paragraph_format.space_after = Pt(4)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE_DARK if style_name == "Heading 1" else NAVY)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header.distance = Inches(0.25)
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("PLASMIT HOSPITAL HMS"), size=8, bold=True, color=BLUE_DARK)
    set_run(hp.add_run("   |   ICU Command Center"), size=8, color=SLATE)

    footer = section.footer
    footer.distance = Inches(0.25)
    fp = footer.paragraphs[0]
    fp.add_run("Confidential - For Internal Leadership Review")
    for run in fp.runs:
        set_run(run, size=7.5, color=SLATE)
    add_page_number(footer.add_paragraph())

    # Cover
    cover = doc.add_table(rows=1, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = cover.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=360, start=300, bottom=360, end=300)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run("PLASMIT HOSPITAL HMS"), size=12, bold=True, color="7DD3FC")
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("ICU Command Center"), size=30, bold=True, color=WHITE, font="Aptos Display")
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(25)
    set_run(p.add_run("Executive Use Case Overview"), size=19, bold=True, color="BAE6FD", font="Aptos Display")
    p = cell.add_paragraph()
    set_run(p.add_run("A high-level view of command, clinical, nursing, diagnostics, tele-ICU, device, intelligence, analytics and governance workflows."), size=10.5, color="E2E8F0")

    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    meta.columns[0].width = Inches(1.65)
    meta.columns[1].width = Inches(5.65)
    metadata = [
        ("Prepared for", "Senior Leadership and Clinical Governance"),
        ("Document type", "High-Level Functional Overview"),
        ("Scope", "ICU Command Center - Command through Administration"),
        ("Date", "22 June 2026"),
    ]
    for row, (label, value) in zip(meta.rows, metadata):
        set_cell_shading(row.cells[0], BLUE_LIGHT)
        set_cell_shading(row.cells[1], WHITE)
        for cell_item in row.cells:
            set_cell_margins(cell_item, top=100, bottom=100)
            set_cell_border(cell_item, bottom={"val": "single", "sz": "5", "color": "CBD5E1"})
        set_run(row.cells[0].paragraphs[0].add_run(label), size=8.5, bold=True, color=BLUE_DARK)
        set_run(row.cells[1].paragraphs[0].add_run(value), size=8.8, color=NAVY)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(note.add_run("Leadership summary: 11 functional groups | 41 user-facing pages | 24x7 ICU workflow coverage"), size=9.2, bold=True, color=BLUE_DARK)
    doc.add_page_break()

    # Executive overview
    add_heading(doc, "1. Executive Overview", 1)
    add_body(doc, "The ICU Command Center is designed as a unified operational and clinical coordination layer for critical care. It brings patient condition, staffing, medication, diagnostics, devices, alerts, escalation and management insight into a role-aware workspace.")
    add_body(doc, "The platform supports faster prioritization and clearer accountability while keeping final clinical decisions with authorized care teams.")

    cards = [
        ("11 Functional Groups", "Command through Administration, organized around real ICU operating responsibilities."),
        ("41 User-Facing Pages", "Each page has a clear user, decision purpose and accountable outcome."),
        ("Patient + Unit View", "Supports bedside review as well as unit and enterprise-level command."),
        ("Closed-Loop Action", "Alert, assignment, SLA, response, handover and audit are connected."),
    ]
    add_two_column_cards(doc, cards)

    add_heading(doc, "Leadership Value", 2)
    for item in [
        "One operational picture across all ICU units and patient beds.",
        "Earlier recognition and escalation of deterioration, delays and device issues.",
        "Clear responsibility through Assigned To, SLA, status and documented outcome.",
        "Stronger medication, nursing, diagnostics and shift-handover discipline.",
        "Auditable workflows and measurable performance for quality improvement.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Primary Personas", 1)
    add_two_column_cards(doc, PERSONAS)

    add_heading(doc, "3. End-to-End ICU Operating Flow", 1)
    flow_table = doc.add_table(rows=1, cols=3)
    flow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_table.autofit = False
    flow_table.columns[0].width = Inches(0.55)
    flow_table.columns[1].width = Inches(2.05)
    flow_table.columns[2].width = Inches(5.0)
    for idx, title in enumerate(("STEP", "STAGE", "SYSTEM SUPPORT")):
        flow_table.rows[0].cells[idx].text = title
    style_header_row(flow_table.rows[0])
    for idx, (step, stage, support) in enumerate(FLOW):
        cells = flow_table.add_row().cells
        for cell_item in cells:
            set_cell_shading(cell_item, WHITE if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell_item)
            set_cell_border(cell_item, bottom={"val": "single", "sz": "4", "color": "CBD5E1"})
        cells[0].text = step
        cells[1].text = stage
        cells[2].text = support
        for cell_item in cells:
            for run in cell_item.paragraphs[0].runs:
                set_run(run, size=8.6, color=SLATE)
        for run in cells[0].paragraphs[0].runs:
            set_run(run, size=9, bold=True, color=BLUE_DARK)
        for run in cells[1].paragraphs[0].runs:
            set_run(run, size=8.8, bold=True, color=NAVY)

    doc.add_page_break()
    add_heading(doc, "4. Module and Page Use Cases", 1)
    add_body(doc, "The following pages represent the approved high-level functional scope. Each page is listed with its primary users and the decision or workflow it supports.")

    for module_index, (module, purpose, rows) in enumerate(MODULES, start=1):
        if module in {"Clinical Workspace", "Diagnostics", "Clinical Intelligence"}:
            doc.add_page_break()
        add_heading(doc, f"4.{module_index} {module}", 2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run(p.add_run(purpose), size=8.8, color=SLATE)
        add_module_table(doc, rows)

    doc.add_page_break()
    add_heading(doc, "5. Governance and Control Principles", 1)
    governance = [
        ("Role-Based Access", "Users see actions appropriate to their role, unit and responsibility."),
        ("Patient Context", "Every clinical action remains linked to the correct patient, bed and encounter."),
        ("Human Verification", "Scores, extracted reports and automated alerts support review; they do not replace clinical judgement."),
        ("Auditability", "Important views, acknowledgements, assignments and changes are retained for review."),
        ("Configurable Rules", "Thresholds, SLAs, escalation paths and operational parameters are governed centrally."),
        ("Downtime Readiness", "Critical workflows require a defined manual fallback when network or device data is unavailable."),
    ]
    add_two_column_cards(doc, governance)

    add_heading(doc, "6. High-Level Integration Landscape", 1)
    integration_table = doc.add_table(rows=1, cols=2)
    integration_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    integration_table.autofit = False
    integration_table.columns[0].width = Inches(2.35)
    integration_table.columns[1].width = Inches(5.25)
    integration_table.rows[0].cells[0].text = "SOURCE / SERVICE"
    integration_table.rows[0].cells[1].text = "ICU COMMAND CENTER USE"
    style_header_row(integration_table.rows[0])
    integrations = [
        ("HMS / EMR / Admission", "Patient identity, encounter, diagnosis, care team, admission and discharge context."),
        ("Bedside Devices / Gateways", "Vitals, ventilator, pump, signal health and device-state visibility."),
        ("LIS / RIS / PACS", "Orders, samples, structured results, imaging reports and critical values."),
        ("Pharmacy / eMAR", "Medication orders, dispensing, due status, administration and safety checks."),
        ("FHIR / HL7 APIs", "Standards-based exchange between clinical and operational systems."),
        ("Identity / Audit Services", "Role access, user traceability and compliance evidence."),
    ]
    for idx, (source, use) in enumerate(integrations):
        cells = integration_table.add_row().cells
        cells[0].text = source
        cells[1].text = use
        for cell_item in cells:
            set_cell_shading(cell_item, WHITE if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell_item)
            set_cell_border(cell_item, bottom={"val": "single", "sz": "4", "color": "CBD5E1"})
            for run in cell_item.paragraphs[0].runs:
                set_run(run, size=8.6, color=SLATE)
        for run in cells[0].paragraphs[0].runs:
            set_run(run, size=8.7, bold=True, color=NAVY)

    add_heading(doc, "7. Recommended Rollout Approach", 1)
    rollout = [
        ("1. Clinical Validation", "Confirm workflows, roles, escalation thresholds and documentation requirements with ICU leadership."),
        ("2. Live Integration", "Connect patient, diagnostics, pharmacy and device data through governed interfaces."),
        ("3. Controlled Pilot", "Activate one ICU unit with trained champions and monitored fallback processes."),
        ("4. Measure Outcomes", "Track response time, task closure, medication compliance, data completeness and user adoption."),
        ("5. Scale by Readiness", "Expand unit-by-unit after clinical, security, performance and governance sign-off."),
    ]
    add_two_column_cards(doc, rollout)

    add_heading(doc, "Management Decision Summary", 2)
    summary_box = doc.add_table(rows=1, cols=1)
    summary_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_cell = summary_box.cell(0, 0)
    set_cell_shading(summary_cell, BLUE_LIGHT)
    set_cell_border(summary_cell, top={"val": "single", "sz": "8", "color": BLUE}, bottom={"val": "single", "sz": "8", "color": BLUE}, start={"val": "single", "sz": "8", "color": BLUE}, end={"val": "single", "sz": "8", "color": BLUE})
    set_cell_margins(summary_cell, top=170, start=190, bottom=170, end=190)
    p = summary_cell.paragraphs[0]
    set_run(p.add_run("The ICU Command Center provides a coherent blueprint for safer, faster and more accountable critical-care operations. The recommended next step is a clinically governed pilot with live data integration and measurable success criteria."), size=10, bold=True, color=NAVY)

    add_body(doc, "Scope note: This document describes the high-level functional intent of the current ICU Command Center design. Live-data readiness, clinical validation, integration certification, security testing and production rollout remain subject to formal approval.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output_path = build_document()
    print(output_path)
